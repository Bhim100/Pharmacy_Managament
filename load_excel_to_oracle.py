import datetime as dt
import os
import sys
from pathlib import Path

try:
    import oracledb
except ImportError as exc:
    raise SystemExit("Missing dependency: oracledb. Run 'pip install -r backend/requirements.txt'.") from exc

try:
    from openpyxl import load_workbook
except ImportError as exc:
    raise SystemExit("Missing dependency: openpyxl. Run 'pip install -r backend/requirements.txt'.") from exc


PROJECT_ROOT = Path(__file__).resolve().parent
DATA_DIR = PROJECT_ROOT / "data"

TABLE_IMPORT_ORDER = [
    {
        "file_name": "01_categories.xlsx",
        "table_name": "categories",
        "date_columns": set(),
    },
    {
        "file_name": "02_suppliers.xlsx",
        "table_name": "suppliers",
        "date_columns": set(),
    },
    {
        "file_name": "03_medicines.xlsx",
        "table_name": "medicines",
        "date_columns": {"created_at"},
    },
    {
        "file_name": "04_batches.xlsx",
        "table_name": "batches",
        "date_columns": {"manufacturing_date", "expiry_date", "received_date"},
    },
    {
        "file_name": "05_customers.xlsx",
        "table_name": "customers",
        "date_columns": {"date_of_birth", "created_at"},
    },
    {
        "file_name": "06_doctors.xlsx",
        "table_name": "doctors",
        "date_columns": set(),
    },
]


def get_connection():
    """Build an Oracle connection from the same env vars used by Flask."""
    user = os.getenv("PHARMACY_DB_USER")
    password = os.getenv("PHARMACY_DB_PASSWORD")
    dsn = os.getenv("PHARMACY_DB_DSN")

    if not user or not password or not dsn:
        raise SystemExit(
            "Missing PHARMACY_DB_USER, PHARMACY_DB_PASSWORD, or PHARMACY_DB_DSN. "
            "Set them in the terminal before running this loader."
        )

    return oracledb.connect(user=user, password=password, dsn=dsn)


def normalize_value(column_name, value, date_columns):
    """Convert Excel cell values into Oracle-friendly Python values."""
    if value is None:
        return None

    if isinstance(value, str):
        cleaned = value.strip()
        if cleaned == "":
            return None
        if column_name in date_columns:
            return dt.datetime.strptime(cleaned, "%Y-%m-%d").date()
        return cleaned

    if column_name in date_columns:
        if isinstance(value, dt.datetime):
            return value.date()
        if isinstance(value, dt.date):
            return value

    return value


def load_rows_from_workbook(workbook_path, date_columns):
    """Read one Excel file and return headers plus cleaned row dictionaries."""
    workbook = load_workbook(workbook_path, data_only=True)
    worksheet = workbook.active

    headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
    rows = []

    for row in worksheet.iter_rows(min_row=2, values_only=True):
        if all(value in (None, "") for value in row):
            continue

        row_payload = {}
        for header, value in zip(headers, row):
            row_payload[header] = normalize_value(header, value, date_columns)
        rows.append(row_payload)

    return headers, rows


def delete_existing_rows(cursor):
    """Delete imported table data in reverse order so foreign keys stay valid."""
    for config in reversed(TABLE_IMPORT_ORDER):
        print(f"Clearing table: {config['table_name']}")
        cursor.execute(f"DELETE FROM {config['table_name']}")


def import_table(cursor, config):
    """Load one workbook into its matching Oracle table."""
    workbook_path = DATA_DIR / config["file_name"]
    headers, rows = load_rows_from_workbook(workbook_path, config["date_columns"])

    if not rows:
        print(f"Skipping {config['table_name']} because no rows were found in {config['file_name']}.")
        return

    column_list = ", ".join(headers)
    bind_list = ", ".join(f":{header}" for header in headers)
    sql = f"INSERT INTO {config['table_name']} ({column_list}) VALUES ({bind_list})"

    # executemany keeps the import fast and avoids writing one INSERT per row by hand.
    cursor.executemany(sql, rows)
    print(f"Loaded {len(rows)} rows into {config['table_name']}.")


def main():
    replace_existing = "--replace" in sys.argv

    if not DATA_DIR.exists():
        raise SystemExit(f"Data folder not found: {DATA_DIR}")

    with get_connection() as connection:
        with connection.cursor() as cursor:
            if replace_existing:
                delete_existing_rows(cursor)

            # Parent tables are imported first so child-table foreign keys can resolve cleanly.
            for config in TABLE_IMPORT_ORDER:
                import_table(cursor, config)

        connection.commit()

    print("Excel import complete.")


if __name__ == "__main__":
    main()
